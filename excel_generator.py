import os
import re
import tempfile
from typing import List, Dict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HEADERS = [
    "Motivo",
    "Data",
    "Horário",
    "Funcionário",
    "Origem",
    "Parada",
    "Destino",
    "Valor da Viagem",
    "Centro de Custo",
    "Data de Registro",
]


def _clean_price_value(val_str: str) -> float:
    """Converte strings como 'R$ 28,96' ou '28.96' para float."""
    if not val_str:
        return 0.0
    try:
        cleaned = re.sub(r"[^\d,\.]", "", val_str)
        if "," in cleaned and "." in cleaned:
            cleaned = cleaned.replace(".", "").replace(",", ".")
        elif "," in cleaned:
            cleaned = cleaned.replace(",", ".")
        return float(cleaned)
    except Exception:
        return 0.0


def create_styled_reimbursement_excel(
    records: List[Dict[str, str]], 
    start_str: str, 
    end_str: str,
    output_dir: str = None
) -> str:
    """
    Gera uma planilha Excel altamente estilizada e profissional (.xlsx) contendo os registros filtrados.
    Retorna o caminho absoluto do arquivo salvo.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Relatório de Viagens"
    ws.views.sheetView[0].showGridLines = True

    # Paleta de Cores (Estilo Dark Slate & Modern Corporate)
    TITLE_BG = "0F172A"       # Azul Escuro Escuro
    SUBTITLE_BG = "1E293B"    # Azul Grafite
    HEADER_BG = "1E293B"      # Cabeçalho da Tabela
    ZEBRA_EVEN = "F8FAFC"     # Linhas pares (Cinza muito claro)
    ZEBRA_ODD = "FFFFFF"      # Linhas ímpares (Branco)
    TOTAL_BG = "E2E8F0"       # Linha de Totais
    CARD_BG = "F1F5F9"        # Card KPI
    FONT_FAMILY = "Segoe UI"

    # EstilosReutilizáveis
    font_title = Font(name=FONT_FAMILY, size=16, bold=True, color="FFFFFF")
    font_subtitle = Font(name=FONT_FAMILY, size=11, italic=True, color="E2E8F0")
    font_kpi_label = Font(name=FONT_FAMILY, size=9, bold=True, color="475569")
    font_kpi_val = Font(name=FONT_FAMILY, size=14, bold=True, color="0F172A")
    font_header = Font(name=FONT_FAMILY, size=10, bold=True, color="FFFFFF")
    font_data = Font(name=FONT_FAMILY, size=10, color="1E293B")
    font_total = Font(name=FONT_FAMILY, size=11, bold=True, color="0F172A")

    fill_title = PatternFill(start_color=TITLE_BG, end_color=TITLE_BG, fill_type="solid")
    fill_subtitle = PatternFill(start_color=SUBTITLE_BG, end_color=SUBTITLE_BG, fill_type="solid")
    fill_header = PatternFill(start_color=HEADER_BG, end_color=HEADER_BG, fill_type="solid")
    fill_even = PatternFill(start_color=ZEBRA_EVEN, end_color=ZEBRA_EVEN, fill_type="solid")
    fill_odd = PatternFill(start_color=ZEBRA_ODD, end_color=ZEBRA_ODD, fill_type="solid")
    fill_total = PatternFill(start_color=TOTAL_BG, end_color=TOTAL_BG, fill_type="solid")
    fill_card = PatternFill(start_color=CARD_BG, end_color=CARD_BG, fill_type="solid")

    thin_border_side = Side(border_style="thin", color="CBD5E1")
    border_data = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    border_total = Border(
        top=Side(border_style="thin", color="0F172A"),
        bottom=Side(border_style="double", color="0F172A"),
        left=thin_border_side,
        right=thin_border_side
    )

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # 1. Título Principal (Linhas 1 e 2)
    ws.merge_cells("A1:J1")
    title_cell = ws["A1"]
    title_cell.value = "RELATÓRIO MENSAL DE VIAGENS"
    title_cell.font = font_title
    title_cell.fill = fill_title
    title_cell.alignment = align_center
    ws.row_dimensions[1].height = 32

    ws.merge_cells("A2:J2")
    sub_cell = ws["A2"]
    sub_cell.value = f"Período de Referência: {start_str} até {end_str}"
    sub_cell.font = font_subtitle
    sub_cell.fill = fill_subtitle
    sub_cell.alignment = align_center
    ws.row_dimensions[2].height = 22

    # 2. Seção de Cards KPI (Linhas 4 e 5)
    total_viagens = len(records)
    total_valor = sum(_clean_price_value(r.get("Valor da Viagem", r.get("valor", ""))) for r in records)

    # Card 1: Total Viagens (B4:C5)
    ws.merge_cells("B4:C4")
    ws["B4"].value = "TOTAL DE VIAGENS"
    ws["B4"].font = font_kpi_label
    ws["B4"].alignment = align_center
    ws["B4"].fill = fill_card

    ws.merge_cells("B5:C5")
    ws["B5"].value = total_viagens
    ws["B5"].font = font_kpi_val
    ws["B5"].alignment = align_center
    ws["B5"].fill = fill_card

    # Card 2: Valor Total (H4:I5)
    ws.merge_cells("H4:I4")
    ws["H4"].value = "VALOR TOTAL REEMBOLSADO"
    ws["H4"].font = font_kpi_label
    ws["H4"].alignment = align_center
    ws["H4"].fill = fill_card

    ws.merge_cells("H5:I5")
    ws["H5"].value = total_valor
    ws["H5"].font = font_kpi_val
    ws["H5"].number_format = 'R$ #,##0.00'
    ws["H5"].alignment = align_center
    ws["H5"].fill = fill_card

    ws.row_dimensions[4].height = 18
    ws.row_dimensions[5].height = 24

    # 3. Cabeçalhos da Tabela (Linha 7)
    header_row_idx = 7
    ws.row_dimensions[header_row_idx].height = 28
    for col_idx, h_text in enumerate(HEADERS, 1):
        cell = ws.cell(row=header_row_idx, column=col_idx, value=h_text)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_data

    # 4. Inserção das Linhas de Dados (Linha 8 em diante)
    current_row = 8
    for r_idx, record in enumerate(records):
        ws.row_dimensions[current_row].height = 22
        row_fill = fill_even if r_idx % 2 == 0 else fill_odd

        # Extração de valores
        motivo = record.get("Motivo", record.get("motivo", ""))
        data = record.get("Data", record.get("data", ""))
        horario = record.get("Horário", record.get("horario", ""))
        funcionario = record.get("Funcionário", record.get("funcionario", ""))
        origem = record.get("Origem", record.get("origem", ""))
        parada = record.get("Parada", record.get("parada", ""))
        destino = record.get("Destino", record.get("destino", ""))
        valor_raw = record.get("Valor da Viagem", record.get("valor", ""))
        valor_float = _clean_price_value(valor_raw)
        centro_custo = record.get("Centro de Custo", record.get("centro_custo", ""))
        data_reg = record.get("Data de Registro", "")

        row_values = [
            (motivo, align_left, "@"),
            (data, align_center, "@"),
            (horario, align_center, "@"),
            (funcionario, align_left, "@"),
            (origem, align_left, "@"),
            (parada, align_left, "@"),
            (destino, align_left, "@"),
            (valor_float, align_right, 'R$ #,##0.00'),
            (centro_custo, align_left, "@"),
            (data_reg, align_center, "@"),
        ]

        for c_idx, (val, alignment, num_fmt) in enumerate(row_values, 1):
            cell = ws.cell(row=current_row, column=c_idx, value=val)
            cell.font = font_data
            cell.fill = row_fill
            cell.alignment = alignment
            cell.border = border_data
            if num_fmt and num_fmt != "@":
                cell.number_format = num_fmt

        current_row += 1

    # 5. Linha de Totalização
    ws.row_dimensions[current_row].height = 26
    for c_idx in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=current_row, column=c_idx)
        cell.font = font_total
        cell.fill = fill_total
        cell.border = border_total

    ws.cell(row=current_row, column=1, value="TOTAL GERAL").alignment = align_left
    sum_col_letter = get_column_letter(8) # Coluna H (Valor)
    if records:
        start_cell = f"{sum_col_letter}8"
        end_cell = f"{sum_col_letter}{current_row - 1}"
        val_cell = ws.cell(row=current_row, column=8, value=f"=SUM({start_cell}:{end_cell})")
    else:
        val_cell = ws.cell(row=current_row, column=8, value=0.0)

    val_cell.number_format = 'R$ #,##0.00'
    val_cell.alignment = align_right

    # 6. Quadrinho de Resumo por Centro de Custo (abaixo da tabela principal)
    from collections import defaultdict

    cc_data = defaultdict(lambda: {"count": 0, "total": 0.0})
    for r in records:
        cc_name = r.get("Centro de Custo", r.get("centro_custo", "")).strip() or "Não Informado"
        v_float = _clean_price_value(r.get("Valor da Viagem", r.get("valor", "")))
        cc_data[cc_name]["count"] += 1
        cc_data[cc_name]["total"] += v_float

    # Ordena pelos centros de custo com maior valor gasto
    sorted_cc = sorted(cc_data.items(), key=lambda x: x[1]["total"], reverse=True)

    cc_box_start = current_row + 3
    ws.merge_cells(f"B{cc_box_start}:E{cc_box_start}")
    cc_title = ws[f"B{cc_box_start}"]
    cc_title.value = "RESUMO DE GASTOS POR CENTRO DE CUSTO"
    cc_title.font = Font(name=FONT_FAMILY, size=11, bold=True, color="FFFFFF")
    cc_title.fill = fill_title
    cc_title.alignment = align_center
    ws.row_dimensions[cc_box_start].height = 26

    cc_headers = [
        ("Centro de Custo", align_left),
        ("Qtd. Viagens", align_center),
        ("Valor Total", align_right),
        ("% do Total", align_center),
    ]

    cc_header_row = cc_box_start + 1
    ws.row_dimensions[cc_header_row].height = 22
    for idx, (h_name, h_align) in enumerate(cc_headers, start=2):
        c = ws.cell(row=cc_header_row, column=idx, value=h_name)
        c.font = font_header
        c.fill = fill_header
        c.alignment = h_align
        c.border = border_data

    cc_curr_row = cc_header_row + 1
    first_cc_row = cc_curr_row
    for idx, (cc_name, info) in enumerate(sorted_cc):
        ws.row_dimensions[cc_curr_row].height = 20
        row_fill = fill_even if idx % 2 == 0 else fill_odd
        pct = (info["total"] / total_valor) if total_valor > 0 else 0.0

        c_name = ws.cell(row=cc_curr_row, column=2, value=cc_name)
        c_name.alignment = align_left
        c_name.font = font_data
        c_name.fill = row_fill
        c_name.border = border_data

        c_cnt = ws.cell(row=cc_curr_row, column=3, value=info["count"])
        c_cnt.alignment = align_center
        c_cnt.font = font_data
        c_cnt.fill = row_fill
        c_cnt.border = border_data

        c_val = ws.cell(row=cc_curr_row, column=4, value=info["total"])
        c_val.alignment = align_right
        c_val.font = font_data
        c_val.fill = row_fill
        c_val.border = border_data
        c_val.number_format = 'R$ #,##0.00'

        c_pct = ws.cell(row=cc_curr_row, column=5, value=pct)
        c_pct.alignment = align_center
        c_pct.font = font_data
        c_pct.fill = row_fill
        c_pct.border = border_data
        c_pct.number_format = '0.0%'

        cc_curr_row += 1

    # Linha de Total do Centro de Custo
    ws.row_dimensions[cc_curr_row].height = 24
    for c_idx in range(2, 6):
        c = ws.cell(row=cc_curr_row, column=c_idx)
        c.font = font_total
        c.fill = fill_total
        c.border = border_total

    ws.cell(row=cc_curr_row, column=2, value="TOTAL").alignment = align_left
    if sorted_cc:
        ws.cell(row=cc_curr_row, column=3, value=f"=SUM(C{first_cc_row}:C{cc_curr_row - 1})").alignment = align_center
        c_sum = ws.cell(row=cc_curr_row, column=4, value=f"=SUM(D{first_cc_row}:D{cc_curr_row - 1})")
        c_sum.number_format = 'R$ #,##0.00'
        c_sum.alignment = align_right
        c_pct_tot = ws.cell(row=cc_curr_row, column=5, value=f"=SUM(E{first_cc_row}:E{cc_curr_row - 1})")
        c_pct_tot.number_format = '0.0%'
        c_pct_tot.alignment = align_center
    else:
        ws.cell(row=cc_curr_row, column=3, value=0).alignment = align_center
        c_sum = ws.cell(row=cc_curr_row, column=4, value=0.0)
        c_sum.number_format = 'R$ #,##0.00'
        c_sum.alignment = align_right
        ws.cell(row=cc_curr_row, column=5, value=0.0).alignment = align_center

    # 7. Ajuste Automático de Largura das Colunas
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            # Ignora células mescladas no título, cards e título do centro de custo
            if cell.row in (1, 2, 4, 5, cc_box_start):
                continue
            if cell.value:
                val_str = str(cell.value)
                max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    # Definir diretório de saída
    if not output_dir:
        output_dir = os.path.join(tempfile.gettempdir(), "uber_exports")
    os.makedirs(output_dir, exist_ok=True)

    safe_start = start_str.replace("/", "-")
    safe_end = end_str.replace("/", "-")
    filename = f"Relatorio_Viagens_{safe_start}_a_{safe_end}.xlsx"
    filepath = os.path.join(output_dir, filename)

    wb.save(filepath)
    return filepath

